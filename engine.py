"""Local-first extraction and safe folder creation. No network calls."""
import csv, json, re, subprocess
from pathlib import Path
from email import policy
from email.parser import BytesParser

SUPPORTED={'.txt','.rtf','.md','.csv','.tsv','.xlsx','.xlsm','.json','.eml'}
STOP=re.compile(r'\b(?:don.t|do not|skip|exclude|omit|not needed|ignore|hold off|no need|without)\b',re.I)
INTRO=re.compile(r'^(?:please\s+)?(?:create|make|set up|we need|folders?\s*(?:for|:)|categories?\s*:|include|add|organize|need)\s*[:\-]?\s*',re.I)
BULLET=re.compile(r'^\s*(?:[-*•]|\d+[.)]|\[[ xX]\])\s*')
EXT=re.compile(r'\s*(?:\([^)]{0,80}\)|\[[^]]{0,80}\])\s*$')
END=re.compile(r'[.!?;]+$')

def read_text(path):
    path=Path(path); suffix=path.suffix.lower()
    if suffix=='.rtf' or path.read_bytes()[:32].lstrip().startswith(b'{\\rtf'):
        p=subprocess.run(['/usr/bin/textutil','-convert','txt','-stdout',str(path)],capture_output=True)
        if p.returncode: raise ValueError(p.stderr.decode(errors='replace'))
        return p.stdout.decode('utf-8-sig',errors='replace')
    if suffix=='.eml':
        msg=BytesParser(policy=policy.default).parsebytes(path.read_bytes())
        body=msg.get_body(preferencelist=('plain',)) if msg.is_multipart() else msg
        return body.get_content() if body else ''
    for enc in ('utf-8-sig','utf-16','mac_roman'):
        try: return path.read_text(encoding=enc)
        except UnicodeError: pass
    raise ValueError('Unsupported text encoding')

def read_table(path):
    path=Path(path)
    if path.suffix.lower() in ('.xlsx','.xlsm'):
        from openpyxl import load_workbook
        wb=load_workbook(path,read_only=True,data_only=True)
        try: return {ws.title:[list(row) for row in ws.iter_rows(values_only=True)] for ws in wb.worksheets}
        finally: wb.close()
    if path.suffix.lower() in ('.csv','.tsv'):
        with path.open(encoding='utf-8-sig',newline='') as f:
            sample=f.read(8192); f.seek(0)
            try: dialect=csv.Sniffer().sniff(sample,delimiters=',;\t')
            except csv.Error: dialect=csv.excel_tab if path.suffix.lower()=='.tsv' else csv.excel
            return {'Data':list(csv.reader(f,dialect))}
    raise ValueError('Not a supported table')

def valid_name(name):
    name=str(name).strip().strip('"“”\'`').strip()
    if not name or name in ('.','..') or len(name)>180 or any(c in name for c in ('/',':','\0','\n','\r')): return None
    if name.startswith('~') or name.startswith('.'): return None
    return name

def dedupe(items):
    out=[]; seen=set()
    for item in items:
        name=valid_name(item)
        if name and name.casefold() not in seen:
            out.append(name);seen.add(name.casefold())
    return out

def extract(text,mode='smart',split_commas=True):
    """Conservative extraction: never turn arbitrary prose into folder names."""
    if mode=='literal':
        return dedupe([part for line in text.splitlines() for part in (line.split(',') if split_commas else [line])]),[],[]
    results=[]; review=[]; excluded=[]; in_list=False
    lines=text.replace('\r','\n').split('\n')
    for raw in lines:
        line=raw.strip()
        if not line:
            continue
        # Quoted messages and signature blocks are not lists.
        if line.startswith('>') or re.match(r'^(?:hi[!, ]|hello[!, ]|thanks\b|regards\b|best[,:]|p\.p\.s\.|p\.s\.|director of|subject:|from:|to:)',line,re.I):
            in_list=False;continue
        bullet=bool(BULLET.match(line))
        content=BULLET.sub('',line).strip()
        if re.search(r'\b(?:don.t|do not|never create|no folder (?:called|named)|skip creating)\b',content,re.I):
            excluded.append(raw.strip());in_list=False;continue
        # Explicit list header is the only context that promotes bullets to names.
        if re.search(r'^(?:for the |for )?[^.!?]{0,65}?(?:folders? for|folders? (?:called|named)|please separate|categories? for)\s*:\s*$',content,re.I):
            in_list=True;continue
        if re.match(r'^(?:a few other|one last thing|speaking of|oh[!,]|also[!,]|basically[!,]|first[!,]|here.s what|for folder names|for the invoices|we also have|the spreadsheet|the actual|please keep|please don.t)',content,re.I):
            in_list=False
        # Explicit named folder requests anywhere in prose.
        m=re.search(r'\b(?:folder (?:called|named)|folder for)\s+["“]?([^"”.,;!?]{1,65})',content,re.I)
        if m:
            candidate=m.group(1).strip()
            candidate=re.split(r'\s+\b(?:but|because|unless|which|so that|for anything|where)\b',candidate,maxsplit=1,flags=re.I)[0].strip()
            if candidate and not re.search(r'\b(?:anything|everyone|all our|each|every)\b',candidate,re.I):results.append(candidate)
            continue
        # "go under X" signals a destination category, not an arbitrary sentence.
        for match in re.finditer(r'\b(?:go(?:es)? under|put (?:them|it|those) under)\s+([A-Z][A-Za-z ]{1,35}?)(?=\s*(?:[,.;]|\b(?:except|unless|but|because|which|and|$)))',content):
            results.append(match.group(1).strip())
        if bullet and in_list:
            item=content.strip().strip('"“”').rstrip('.;')
            if len(item)<=65 and len(item.split())<=10 and not re.search(r'\b(?:should|would|could|please|unless|except|because|which|when|if|don.t|do not)\b',item,re.I):
                results.append(item)
            else:review.append(content)
            continue
        # Isolated standalone names are supported, but only in an otherwise simple list.
        if len(lines)<=15 and not any(re.search(r'[.!?]',x) for x in lines if x.strip()):
            if len(content.split())<=6 and len(content)<=60 and not re.search(r'\b(?:also|here|thinking|need|want|should|would|please|because|under|folder|anything)\b',content,re.I):
                results.extend(content.split(',') if split_commas else [content]);continue
        if re.search(r'\b(?:folder|categor|organize|sort|under|except|unless|if |should|would|need|anything|involving)\b',content,re.I):
            review.append(content)
    return dedupe(results),dedupe(review),excluded

def from_json(path):
    obj=json.loads(Path(path).read_text(encoding='utf-8-sig'))
    if isinstance(obj,list):
        return [x if isinstance(x,str) else x.get('name','') if isinstance(x,dict) else '' for x in obj]
    if isinstance(obj,dict):
        for k in ('folders','names','items','categories'):
            if isinstance(obj.get(k),list): return [x if isinstance(x,str) else x.get('name','') if isinstance(x,dict) else '' for x in obj[k]]
    raise ValueError('JSON must contain a list or a folders/names/items/categories array')

def plan(names,destination,prefix='',suffix='',number=False,start=1,case='Original',replace_from='',replace_to=''):
    dest=Path(destination)
    planned=[];seen=set()
    for idx,original in enumerate(names):
        name=original.replace(replace_from,replace_to) if replace_from else original
        if case=='Title Case':name=name.title()
        elif case=='UPPERCASE':name=name.upper()
        elif case=='lowercase':name=name.lower()
        name=f'{prefix}{name}{suffix}'
        if number:name=f'{start+idx:02d} - {name}'
        name=valid_name(name)
        if not name or name.casefold() in seen:continue
        seen.add(name.casefold())
        target=dest/name
        planned.append((name,'Exists' if target.exists() else 'Ready'))
    return planned

def create(names,destination):
    dest=Path(destination)
    if not dest.is_dir():raise ValueError('Destination folder does not exist')
    made=[];skipped=[];errors=[]
    for name in dedupe(names):
        target=dest/name
        if target.exists():skipped.append(name);continue
        try:target.mkdir();made.append(str(target))
        except OSError as exc:errors.append(f'{name}: {exc}')
    return made,skipped,errors

def undo(paths):
    removed=[];failed=[]
    for p in reversed(paths):
        try:Path(p).rmdir();removed.append(p) # Never remove nonempty folders.
        except OSError as exc:failed.append(f'{p}: {exc}')
    return removed,failed


# Folder paths are always relative to the chosen destination. No traversal or
# absolute paths are accepted, including paths supplied by imported files.
def valid_folder_path(value):
    raw=str(value).strip().replace('\\','/')
    if not raw or raw.startswith('/') or raw.endswith('/') or '\x00' in raw:
        return None
    parts=raw.split('/')
    if any(part in ('.','..') or not valid_name(part) for part in parts):
        return None
    return '/'.join(valid_name(part) for part in parts)

def sequence(prefix='shot-', start=1, end=10, digits=2, suffix='', step=1):
    if not isinstance(start,int) or not isinstance(end,int) or not isinstance(digits,int) or not isinstance(step,int):
        raise ValueError('Sequence settings must be whole numbers')
    if step<1 or digits<1 or digits>8 or end<start or (end-start)//step>=1000:
        raise ValueError('Choose an ascending range of at most 1,000 folders')
    names=[f'{prefix}{n:0{digits}d}{suffix}' for n in range(start,end+1,step)]
    if any(valid_folder_path(name) is None for name in names):
        raise ValueError('The sequence contains an invalid folder name')
    return names

def interpret_sequence(text):
    # Accepts explicit numeric ranges, never guesses from unrelated prose.
    m=re.search(r'([A-Za-z][\w-]{0,55}?)(\d{1,6})\s*(?:through|to|[-–])\s*(?:([A-Za-z][\w-]{0,55}?))?(\d{1,6})\b',text,re.I)
    if not m:return None
    prefix=m.group(1).strip(' \"\'')
    other=(m.group(3) or '').strip(' \"\'')
    if other and other.casefold()!=prefix.casefold():return None
    return sequence(prefix,int(m.group(2)),int(m.group(4)),max(len(m.group(2)),len(m.group(4))))

def plan_paths(names,destination):
    root=Path(destination)
    result=[];seen=set()
    for value in names:
        name=valid_folder_path(value)
        if not name or name.casefold() in seen:continue
        seen.add(name.casefold())
        target=root.joinpath(*name.split('/'))
        result.append((name,'Exists' if target.exists() else 'Ready'))
    return result

def create_paths(names,destination):
    root=Path(destination)
    if not root.is_dir():raise ValueError('Destination folder does not exist')
    created=[];skipped=[];errors=[]
    for name,status in plan_paths(names,root):
        target=root.joinpath(*name.split('/'))
        if status=='Exists':skipped.append(name);continue
        try:
            missing=[];cur=target
            while cur!=root and not cur.exists():missing.append(cur);cur=cur.parent
            for folder in reversed(missing):
                folder.mkdir()
                created.append(str(folder))
        except OSError as exc:errors.append(f'{name}: {exc}')
    return created,skipped,errors
